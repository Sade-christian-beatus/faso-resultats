import io

import openpyxl
import pytest
from httpx import AsyncClient
from PIL import Image, ImageDraw, ImageFont
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Resultat

_FONT = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf", 20)


def _construire_pdf_scan_fonction_publique(chemin) -> None:
    image = Image.new("RGB", (1400, 200), "white")
    dessin = ImageDraw.Draw(image)
    dessin.text((30, 20), "ADMISSIBLES", font=_FONT, fill="black")
    dessin.text(
        (30, 60),
        "1° BAYALA JEAN-CLAUDE 000015-120-03 B14863543 15/02/01",
        font=_FONT,
        fill="black",
    )
    dessin.text((30, 100), "Arrete la presente liste a 1 admissible.", font=_FONT, fill="black")
    image.save(str(chemin), "PDF")


def _construire_xlsx(lignes: list[list]) -> bytes:
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.append(
        ["Numéro PV", "Jury", "Nom", "Prénom", "Date de naissance", "Décision", "Moyenne"]
    )
    for ligne in lignes:
        feuille.append(ligne)
    buffer = io.BytesIO()
    classeur.save(buffer)
    return buffer.getvalue()


async def _creer_examen(client: AsyncClient, admin_headers: dict) -> str:
    response = await client.post(
        "/api/v1/admin/exams",
        json={"type_examen": "BAC", "annee": 2026, "libelle": "BAC 2026"},
        headers=admin_headers,
    )
    return response.json()["id"]


@pytest.mark.asyncio
async def test_upload_excel_valide_cree_apercu_sans_erreur(
    client: AsyncClient, admin_headers: dict
) -> None:
    examen_id = await _creer_examen(client, admin_headers)
    contenu = _construire_xlsx([["001", "Ouaga 1", "Traore", "Awa", "12/05/2008", "Admis", 13.45]])

    response = await client.post(
        "/api/v1/admin/ingestions",
        data={"examen_id": examen_id, "type_fichier": "EXCEL"},
        files={
            "file": (
                "resultats.xlsx",
                contenu,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=admin_headers,
    )

    assert response.status_code == 201
    body = response.json()
    assert body["statut"] == "PREVISUALISATION"
    assert body["nombre_lignes_detectees"] == 1
    assert body["nombre_erreurs"] == 0
    assert body["lignes"][0]["donnees"]["numero_pv"] == "001"


@pytest.mark.asyncio
async def test_upload_rejette_extension_incompatible(
    client: AsyncClient, admin_headers: dict
) -> None:
    examen_id = await _creer_examen(client, admin_headers)

    response = await client.post(
        "/api/v1/admin/ingestions",
        data={"examen_id": examen_id, "type_fichier": "EXCEL"},
        files={"file": ("resultats.pdf", b"%PDF-1.4", "application/pdf")},
        headers=admin_headers,
    )

    assert response.status_code == 422


@pytest.mark.asyncio
async def test_upload_examen_introuvable(client: AsyncClient, admin_headers: dict) -> None:
    contenu = _construire_xlsx([])

    response = await client.post(
        "/api/v1/admin/ingestions",
        data={"examen_id": "00000000-0000-0000-0000-000000000000", "type_fichier": "EXCEL"},
        files={
            "file": (
                "resultats.xlsx",
                contenu,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=admin_headers,
    )

    assert response.status_code == 404


@pytest.mark.asyncio
async def test_publish_bloque_si_erreurs_restantes(
    client: AsyncClient, admin_headers: dict
) -> None:
    examen_id = await _creer_examen(client, admin_headers)
    # Ligne incomplète : nom manquant.
    contenu = _construire_xlsx([["001", "Ouaga 1", None, "Awa", None, "Admis", None]])

    upload = await client.post(
        "/api/v1/admin/ingestions",
        data={"examen_id": examen_id, "type_fichier": "EXCEL"},
        files={
            "file": (
                "resultats.xlsx",
                contenu,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=admin_headers,
    )
    ingestion_id = upload.json()["id"]

    response = await client.post(
        f"/api/v1/admin/ingestions/{ingestion_id}/publish", headers=admin_headers
    )

    assert response.status_code == 400


@pytest.mark.asyncio
async def test_correction_puis_publication_cree_les_resultats(
    client: AsyncClient, admin_headers: dict, db_session: AsyncSession
) -> None:
    examen_id = await _creer_examen(client, admin_headers)
    contenu = _construire_xlsx([["001", "Ouaga 1", None, "Awa", None, "Admis", None]])

    upload = await client.post(
        "/api/v1/admin/ingestions",
        data={"examen_id": examen_id, "type_fichier": "EXCEL"},
        files={
            "file": (
                "resultats.xlsx",
                contenu,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=admin_headers,
    )
    ingestion = upload.json()

    ligne_corrigee = ingestion["lignes"][0]
    ligne_corrigee["donnees"]["nom"] = "Traore"
    ligne_corrigee["erreurs"] = []

    correction = await client.patch(
        f"/api/v1/admin/ingestions/{ingestion['id']}",
        json={"lignes": [ligne_corrigee]},
        headers=admin_headers,
    )
    assert correction.status_code == 200
    assert correction.json()["nombre_erreurs"] == 0

    publish = await client.post(
        f"/api/v1/admin/ingestions/{ingestion['id']}/publish", headers=admin_headers
    )

    assert publish.status_code == 200
    assert publish.json()["statut"] == "PUBLIEE"

    resultats = (await db_session.execute(select(Resultat))).scalars().all()
    assert len(resultats) == 1
    assert resultats[0].nom == "Traore"
    assert resultats[0].ingestion_id is not None


@pytest.mark.asyncio
async def test_correction_ignore_les_erreurs_envoyees_par_le_client(
    client: AsyncClient, admin_headers: dict
) -> None:
    """Régression (audit 2026-08-17) : le serveur ne doit jamais faire confiance au
    champ `erreurs` envoyé par le client. Un payload qui prétend "erreurs: []" alors
    que `donnees` est toujours incomplet (nom manquant) doit rester bloqué."""
    examen_id = await _creer_examen(client, admin_headers)
    contenu = _construire_xlsx([["001", "Ouaga 1", None, "Awa", None, "Admis", None]])

    upload = await client.post(
        "/api/v1/admin/ingestions",
        data={"examen_id": examen_id, "type_fichier": "EXCEL"},
        files={
            "file": (
                "resultats.xlsx",
                contenu,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=admin_headers,
    )
    ingestion = upload.json()

    ligne_falsifiee = ingestion["lignes"][0]
    # `nom` reste vide, mais le client affirme mensongèrement qu'il n'y a plus d'erreur.
    ligne_falsifiee["erreurs"] = []

    correction = await client.patch(
        f"/api/v1/admin/ingestions/{ingestion['id']}",
        json={"lignes": [ligne_falsifiee]},
        headers=admin_headers,
    )
    assert correction.status_code == 200
    assert correction.json()["nombre_erreurs"] == 1
    assert "nom manquant" in correction.json()["lignes"][0]["erreurs"]

    publish = await client.post(
        f"/api/v1/admin/ingestions/{ingestion['id']}/publish", headers=admin_headers
    )
    assert publish.status_code == 400


@pytest.mark.asyncio
async def test_reject_ingestion(client: AsyncClient, admin_headers: dict) -> None:
    examen_id = await _creer_examen(client, admin_headers)
    contenu = _construire_xlsx([["001", "Ouaga 1", "Traore", "Awa", None, "Admis", None]])

    upload = await client.post(
        "/api/v1/admin/ingestions",
        data={"examen_id": examen_id, "type_fichier": "EXCEL"},
        files={
            "file": (
                "resultats.xlsx",
                contenu,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=admin_headers,
    )
    ingestion_id = upload.json()["id"]

    response = await client.post(
        f"/api/v1/admin/ingestions/{ingestion_id}/reject", headers=admin_headers
    )

    assert response.status_code == 200
    assert response.json()["statut"] == "REJETEE"


@pytest.mark.asyncio
async def test_ingestions_require_auth(client: AsyncClient) -> None:
    response = await client.get("/api/v1/admin/ingestions")

    assert response.status_code in (401, 403)


@pytest.mark.asyncio
async def test_upload_signale_les_colonnes_non_reconnues(
    client: AsyncClient, admin_headers: dict
) -> None:
    examen_id = await _creer_examen(client, admin_headers)
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.append(["Numéro PV", "Jury", "Nom", "Prénom", "Décision", "Adresse"])
    feuille.append(["001", "Ouaga 1", "Traore", "Awa", "Admis", "Secteur 15"])
    buffer = io.BytesIO()
    classeur.save(buffer)

    response = await client.post(
        "/api/v1/admin/ingestions",
        data={"examen_id": examen_id, "type_fichier": "EXCEL"},
        files={
            "file": (
                "resultats.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=admin_headers,
    )

    body = response.json()
    assert body["nombre_erreurs"] == 0
    assert "Adresse" in body["erreurs_fichier"][0]


@pytest.mark.asyncio
async def test_publish_exige_la_confirmation_des_ecarts_du_fichier(
    client: AsyncClient, admin_headers: dict
) -> None:
    """A file-level warning (here an unrecognised column; for a scanned list, an OCR
    count gap) is not tied to a row: publishing it must be an explicit human decision,
    recorded in the audit log, never a silent one."""
    examen_id = await _creer_examen(client, admin_headers)
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.append(["Numéro PV", "Jury", "Nom", "Prénom", "Décision", "Adresse"])
    feuille.append(["001", "Ouaga 1", "Traore", "Awa", "Admis", "Secteur 15"])
    buffer = io.BytesIO()
    classeur.save(buffer)
    ingestion = (
        await client.post(
            "/api/v1/admin/ingestions",
            data={"examen_id": examen_id, "type_fichier": "EXCEL"},
            files={
                "file": (
                    "resultats.xlsx",
                    buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers=admin_headers,
        )
    ).json()
    assert ingestion["erreurs_fichier"]

    refus = await client.post(
        f"/api/v1/admin/ingestions/{ingestion['id']}/publish", headers=admin_headers
    )
    assert refus.status_code == 409
    assert "écarts à vérifier" in refus.json()["detail"]
    assert "Adresse" in refus.json()["detail"]

    publication = await client.post(
        f"/api/v1/admin/ingestions/{ingestion['id']}/publish",
        params={"confirmer_ecarts": "true"},
        headers=admin_headers,
    )
    assert publication.status_code == 200
    assert publication.json()["statut"] == "PUBLIEE"


@pytest.mark.asyncio
async def test_download_template_requires_auth(client: AsyncClient) -> None:
    response = await client.get("/api/v1/admin/ingestions/template")

    assert response.status_code in (401, 403)


@pytest.mark.asyncio
async def test_download_template_renvoie_un_classeur_excel(
    client: AsyncClient, admin_headers: dict
) -> None:
    response = await client.get("/api/v1/admin/ingestions/template", headers=admin_headers)

    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["content-disposition"]

    classeur = openpyxl.load_workbook(io.BytesIO(response.content))
    feuille = classeur.active
    entetes = [cellule.value for cellule in feuille[1]]
    assert "Numéro PV" in entetes
    assert "Nom" in entetes


@pytest.mark.asyncio
async def test_list_ingestions_filtered_by_examen(client: AsyncClient, admin_headers: dict) -> None:
    examen_id = await _creer_examen(client, admin_headers)
    autre_examen_id = await _creer_examen(client, admin_headers)
    contenu = _construire_xlsx([["001", "Ouaga 1", "Traore", "Awa", None, "Admis", None]])

    await client.post(
        "/api/v1/admin/ingestions",
        data={"examen_id": examen_id, "type_fichier": "EXCEL"},
        files={
            "file": (
                "resultats.xlsx",
                contenu,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=admin_headers,
    )

    response = await client.get(
        "/api/v1/admin/ingestions", params={"examen_id": autre_examen_id}, headers=admin_headers
    )

    assert response.status_code == 200
    assert response.json() == []


@pytest.mark.asyncio
async def test_get_ingestion_introuvable(client: AsyncClient, admin_headers: dict) -> None:
    response = await client.get(
        "/api/v1/admin/ingestions/00000000-0000-0000-0000-000000000000",
        headers=admin_headers,
    )

    assert response.status_code == 404


@pytest.mark.asyncio
async def test_upload_concours_direct_avec_decision_par_defaut_et_correction(
    client: AsyncClient, admin_headers: dict, db_session: AsyncSession
) -> None:
    """Bout en bout sur la structure réelle d'un concours direct (Assistants des
    Douanes) : nom+prénom combinés, decision_par_defaut, N°CNIB."""
    examen_response = await client.post(
        "/api/v1/admin/exams",
        json={
            "type_examen": "CONCOURS_DIRECT",
            "annee": 2026,
            "libelle": "Assistants des Douanes 2026",
        },
        headers=admin_headers,
    )
    examen_id = examen_response.json()["id"]

    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.append(
        ["N°", "NOM ET PRENOM(s)", "RECEPISSE-CODE-CENTRE", "N°CNIB", "DATE NAIS.", "CENTRE"]
    )
    feuille.append([23, "BAYALA JEAN-CLAUDE", "005924-002-06", "B14863543", "15/02/01", "Koudo"])
    buffer = io.BytesIO()
    classeur.save(buffer)

    upload = await client.post(
        "/api/v1/admin/ingestions",
        data={
            "examen_id": examen_id,
            "type_fichier": "EXCEL",
            "decision_par_defaut": "ADMISSIBLE",
        },
        files={
            "file": (
                "douanes.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=admin_headers,
    )
    ingestion = upload.json()
    assert ingestion["nombre_erreurs"] == 1  # prenom manquant, à corriger

    ligne = ingestion["lignes"][0]
    assert ligne["donnees"]["decision"] == "ADMISSIBLE"
    assert ligne["donnees"]["numero_cnib"] == "B14863543"

    ligne["donnees"]["nom"] = "BAYALA"
    ligne["donnees"]["prenom"] = "JEAN-CLAUDE"
    ligne["erreurs"] = []
    correction = await client.patch(
        f"/api/v1/admin/ingestions/{ingestion['id']}",
        json={"lignes": [ligne]},
        headers=admin_headers,
    )
    assert correction.json()["nombre_erreurs"] == 0

    publish = await client.post(
        f"/api/v1/admin/ingestions/{ingestion['id']}/publish", headers=admin_headers
    )
    assert publish.status_code == 200

    resultat = (await db_session.execute(select(Resultat))).scalar_one()
    assert resultat.numero_cnib == "B14863543"
    assert resultat.decision == "ADMISSIBLE"
    assert resultat.nom == "BAYALA"
    assert resultat.prenom == "JEAN-CLAUDE"


@pytest.mark.asyncio
async def test_upload_communique_scanne_fonction_publique_bout_en_bout(
    client: AsyncClient, admin_headers: dict, db_session: AsyncSession, tmp_path
) -> None:
    """Bout en bout sur un communiqué scanné de la Fonction publique : upload en type
    PDF (sans préciser qu'il s'agit d'un scan — détection automatique), extraction
    rang/récépissé/code concours/code centre/CNIB, correction du prénom, publication."""
    examen_response = await client.post(
        "/api/v1/admin/exams",
        json={
            "type_examen": "CONCOURS_DIRECT",
            "annee": 2026,
            "libelle": "Chirurgiens-Dentistes 2026",
        },
        headers=admin_headers,
    )
    examen_id = examen_response.json()["id"]

    chemin_pdf = tmp_path / "communique.pdf"
    _construire_pdf_scan_fonction_publique(chemin_pdf)

    upload = await client.post(
        "/api/v1/admin/ingestions",
        data={"examen_id": examen_id, "type_fichier": "PDF"},
        files={"file": ("communique.pdf", chemin_pdf.read_bytes(), "application/pdf")},
        headers=admin_headers,
    )
    ingestion = upload.json()
    assert ingestion["nombre_erreurs"] == 1  # prenom manquant, à corriger

    ligne = ingestion["lignes"][0]
    assert ligne["donnees"]["numero_recepisse"] == "000015"
    assert ligne["donnees"]["code_concours"] == "120"
    assert ligne["donnees"]["code_centre"] == "03"
    assert ligne["donnees"]["rang_numerique"] == 1
    assert ligne["donnees"]["decision"] == "ADMISSIBLE"

    ligne["donnees"]["nom"] = "BAYALA"
    ligne["donnees"]["prenom"] = "JEAN-CLAUDE"
    ligne["erreurs"] = []
    correction = await client.patch(
        f"/api/v1/admin/ingestions/{ingestion['id']}",
        json={"lignes": [ligne]},
        headers=admin_headers,
    )
    assert correction.json()["nombre_erreurs"] == 0

    publish = await client.post(
        f"/api/v1/admin/ingestions/{ingestion['id']}/publish", headers=admin_headers
    )
    assert publish.status_code == 200

    resultat = (await db_session.execute(select(Resultat))).scalar_one()
    assert resultat.numero_pv == "000015"
    assert resultat.numero_recepisse == "000015"
    assert resultat.code_concours == "120"
    assert resultat.code_centre == "03"
    assert resultat.rang_numerique == 1
    assert resultat.rang_affiche == "1°"
    assert resultat.numero_cnib == "B14863543"
