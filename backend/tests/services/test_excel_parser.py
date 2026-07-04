import openpyxl
import pytest

from app.services.ingestion.excel_parser import parser_excel


@pytest.fixture
def fichier_excel(tmp_path):
    def _creer(lignes: list[list]) -> str:
        classeur = openpyxl.Workbook()
        feuille = classeur.active
        feuille.append(
            ["Numéro PV", "Jury", "Nom", "Prénom", "Date de naissance", "Décision", "Moyenne"]
        )
        for ligne in lignes:
            feuille.append(ligne)
        chemin = tmp_path / "resultats.xlsx"
        classeur.save(chemin)
        return str(chemin)

    return _creer


def test_parser_excel_extrait_les_lignes_valides(fichier_excel) -> None:
    chemin = fichier_excel(
        [
            ["001", "Ouaga 1", "Traore", "Awa", "12/05/2008", "Admis", 13.45],
            ["002", "Ouaga 1", "Kabore", "Issa", None, "Ajourne", None],
        ]
    )

    resultat = parser_excel(chemin)

    assert resultat.nombre_lignes == 2
    assert resultat.nombre_erreurs == 0
    assert resultat.lignes[0].donnees["numero_pv"] == "001"
    assert resultat.lignes[0].donnees["moyenne"] == 13.45
    assert resultat.lignes[1].donnees["decision"] == "AJOURNE"


def test_parser_excel_ignore_les_lignes_vides(fichier_excel) -> None:
    chemin = fichier_excel(
        [
            ["001", "Ouaga 1", "Traore", "Awa", "12/05/2008", "Admis", 13.45],
            [None, None, None, None, None, None, None],
        ]
    )

    resultat = parser_excel(chemin)

    assert resultat.nombre_lignes == 1


def test_parser_excel_signale_ligne_incomplete(fichier_excel) -> None:
    chemin = fichier_excel([["001", "Ouaga 1", None, "Awa", None, "Admis", None]])

    resultat = parser_excel(chemin)

    assert resultat.nombre_erreurs == 1
    assert "nom manquant" in resultat.lignes[0].erreurs


def test_parser_excel_fichier_vide(tmp_path) -> None:
    classeur = openpyxl.Workbook()
    chemin = tmp_path / "vide.xlsx"
    classeur.save(chemin)

    resultat = parser_excel(str(chemin))

    assert resultat.nombre_lignes == 0
    assert resultat.erreurs_fichier == ["Fichier vide"]


def test_parser_excel_document_concours_direct_avec_decision_par_defaut(tmp_path) -> None:
    """Reproduit la structure réelle d'une liste d'admissibilité à un concours
    direct (Assistants des Douanes, OCECOS/AGRE) : nom+prénom combinés, pas de
    colonne décision, un champ N°CNIB inconnu du modèle scolaire classique."""
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.append(
        ["N°", "NOM ET PRENOM(s)", "RECEPISSE-CODE-CENTRE", "N°CNIB", "DATE NAIS.", "CENTRE"]
    )
    feuille.append([23, "BAYALA JEAN-CLAUDE", "005924-002-06", "B14863543", "15/02/01", "Koudo"])
    chemin = tmp_path / "douanes.xlsx"
    classeur.save(chemin)

    resultat = parser_excel(str(chemin), decision_par_defaut="ADMISSIBLE")

    assert resultat.nombre_lignes == 1
    ligne = resultat.lignes[0]
    assert ligne.donnees["numero_pv"] == "005924-002-06"
    assert ligne.donnees["jury"] == "Koudo"
    assert ligne.donnees["nom"] == "BAYALA JEAN-CLAUDE"
    assert ligne.donnees["numero_cnib"] == "B14863543"
    assert ligne.donnees["decision"] == "ADMISSIBLE"
    # Année à 2 chiffres, format courant sur ce type de document.
    assert ligne.donnees["date_naissance"] == "2001-02-15"
    # Prénom non séparé automatiquement : reste à corriger manuellement.
    assert "prenom manquant" in ligne.erreurs
    assert not any("illisible" in erreur for erreur in ligne.erreurs)
